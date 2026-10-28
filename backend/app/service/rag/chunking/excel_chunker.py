'''Excel 解析分块器（P1）：按 Sheet 行组分块，每组携带表头，不做语义切分'''
import logging

from langchain_core.documents import Document
from openpyxl import load_workbook

from app.utils.config_handler import retrieval_conf

logger = logging.getLogger(__name__)

_DEFAULT_ROWS_PER_CHUNK = 20


def chunk_excel(file_path: str) -> list[Document]:
    '''每个 Sheet 按行切块：块内首两行给 Sheet 信息与表头，数据行带原始行号'''
    docs: list[Document] = []
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        for ws in workbook.worksheets:
            # 收集非空行，保留原始行号（从 1 计）；首条非空行作表头
            rows = []
            for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = ["" if v is None else str(v).strip() for v in row]
                if any(values):
                    rows.append((row_no, values))
            if not rows:
                continue
            header_no, header = rows[0]
            data_rows = rows[1:]
            if not data_rows:
                # 只有表头的 Sheet：表头本身作为一个块
                docs.append(Document(
                    page_content=f"【Sheet：{ws.title}｜第{header_no}行】\n表头：" + " | ".join(header),
                    metadata={"sheet_name": ws.title, "row_start": header_no, "row_end": header_no},
                ))
                continue
            step = int(retrieval_conf.get("chunking", {}).get("excel_rows_per_chunk", _DEFAULT_ROWS_PER_CHUNK))
            for i in range(0, len(data_rows), step):
                group = data_rows[i:i + step]
                start_no = group[0][0]
                end_no = group[-1][0]
                lines = [f"【Sheet：{ws.title}｜第{start_no}~{end_no}行】", "表头：" + " | ".join(header)]
                for row_no, values in group:
                    pairs = [f"{h}：{v}" for h, v in zip(header, values) if v]
                    if pairs:
                        lines.append(f"第{row_no}行 " + "；".join(pairs))
                    else:
                        lines.append(f"第{row_no}行")
                docs.append(Document(
                    page_content="\n".join(lines),
                    metadata={"sheet_name": ws.title, "row_start": start_no, "row_end": end_no},
                ))
    finally:
        workbook.close()
    return docs
