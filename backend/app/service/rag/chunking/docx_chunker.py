'''DOCX 解析分块器（P1）

- 按 Heading 1~3 样式构建章节树，标题层级入栈弹栈维护 section_path，标题作为单元正文开头
- 表格独立成块并序列化为文本（首行表头，数据行按“列名：值”展开），元数据带全局 table_index
- .doc（旧格式）无法被 python-docx 解析，抛出明确转换提示
'''
import logging

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.documents import Document

logger = logging.getLogger(__name__)

# 兼容英文/中文 Word 内置标题样式名
_HEADING_STYLES = {
    "Heading 1": 1, "Heading 2": 2, "Heading 3": 3,
    "标题 1": 1, "标题 2": 2, "标题 3": 3,
}


def chunk_docx(file_path: str) -> list[Document]:
    try:
        d = docx.Document(file_path)
    except Exception as e:
        raise ValueError(
            f"文档解析失败（{file_path}）：旧版 .doc 格式不受支持，请另存为 .docx 后重新上传。原始错误：{e}"
        ) from e

    docs: list[Document] = []
    stack: list[tuple[int, str]] = []  # [(level, title)]
    buf: list[str] = []
    table_index = 0

    def flush():
        nonlocal buf
        if not buf:
            return
        meta = {}
        if stack:
            meta["section_path"] = " > ".join(t for _, t in stack)
        docs.append(Document(page_content="\n".join(buf), metadata=meta))
        buf = []

    # 按文档流顺序遍历段落与表格，保证阅读顺序
    for child in d.element.body.iterchildren():
        tag = child.tag
        if tag.endswith("}p"):
            para = Paragraph(child, d)
            text = para.text.strip()
            if not text:
                continue
            level = _HEADING_STYLES.get(getattr(para.style, "name", "") or "")
            if level is not None:
                flush()
                while stack and stack[-1][0] >= level:
                    stack.pop()
                stack.append((level, text))
                buf.append(text)
            else:
                buf.append(text)
        elif tag.endswith("}tbl"):
            serialized = _serialize_table(Table(child, d))
            if not serialized:
                continue
            flush()  # 先落盘表格前的正文，保持文档顺序
            table_index += 1
            meta = {"table_index": table_index}
            if stack:
                meta["section_path"] = " > ".join(t for _, t in stack)
            docs.append(Document(page_content=serialized, metadata=meta))
    flush()
    return docs


def _serialize_table(table: Table) -> str:
    '''表格序列化为文本：首行作表头，数据行按“列名：值”展开'''
    rows = [[(cell.text or "").strip() for cell in row.cells] for row in table.rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    header, data = rows[0], rows[1:]
    lines = ["表格字段：" + " | ".join(header)]
    for r in data:
        pairs = [f"{h}：{v}" for h, v in zip(header, r)]
        lines.append("；".join(pairs))
    return "\n".join(lines)
